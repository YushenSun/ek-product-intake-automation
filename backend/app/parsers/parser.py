from __future__ import annotations

import csv
import io
from dataclasses import dataclass
from pathlib import Path


class ParseError(ValueError):
    pass


@dataclass(frozen=True)
class ParsedDocument:
    text: str
    source_type: str


def parse_text(text: str, source_type: str = "txt") -> ParsedDocument:
    cleaned = text.strip()
    if not cleaned:
        raise ParseError("The submission is empty.")
    return ParsedDocument(text=cleaned, source_type=source_type)


def _rows_to_text(rows: list[list[object]]) -> str:
    if not rows or not rows[0]:
        raise ParseError("Spreadsheet has no usable rows.")
    headers = [str(cell).strip() for cell in rows[0]]
    values = next((row for row in rows[1:] if any(cell not in (None, "") for cell in row)), None)
    if values is None:
        raise ParseError("Spreadsheet contains headers but no values.")
    return "\n".join(f"{header}: {value}" for header, value in zip(headers, values) if header and value not in (None, ""))


def parse_file(filename: str, content: bytes) -> ParsedDocument:
    if not content:
        raise ParseError("The uploaded file is empty.")
    extension = Path(filename).suffix.lower()
    try:
        if extension in {".txt", ".md"}:
            return parse_text(content.decode("utf-8"), "txt")
        if extension == ".csv":
            decoded = content.decode("utf-8-sig")
            rows = list(csv.reader(io.StringIO(decoded)))
            return parse_text(_rows_to_text(rows), "csv")
        if extension == ".xlsx":
            from openpyxl import load_workbook
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            sheet = workbook.active
            rows = [list(row) for row in sheet.iter_rows(values_only=True)]
            return parse_text(_rows_to_text(rows), "xlsx")
        if extension == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(content))
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
            return parse_text(text, "pdf")
    except UnicodeDecodeError as exc:
        raise ParseError("Text file is not valid UTF-8.") from exc
    except Exception as exc:
        raise ParseError(f"Could not parse {extension or 'file'}: {exc}") from exc
    raise ParseError(f"Unsupported file type: {extension or 'no extension'}.")
